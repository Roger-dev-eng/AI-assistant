"""Generate plain DOCX files in the Windows Known Folder Desktop without overwriting files."""

import ctypes
import io
import os
import re
from pathlib import Path
from urllib.parse import urlsplit
from uuid import UUID, uuid4

from assistant.security.urls import public_https_url


def desktop_directory():
    if os.name != "nt":
        raise RuntimeError("A Área de Trabalho automática requer Windows.")
    from ctypes import wintypes

    folder_id = ctypes.create_string_buffer(UUID("B4BFCC3A-DB2C-424C-B029-7FE99A87C641").bytes_le)
    path = wintypes.LPWSTR()
    get_folder = ctypes.windll.shell32.SHGetKnownFolderPath
    get_folder.argtypes = [ctypes.c_void_p, wintypes.DWORD, ctypes.c_void_p, ctypes.POINTER(wintypes.LPWSTR)]
    get_folder.restype = ctypes.c_long
    if get_folder(ctypes.byref(folder_id), 0, None, ctypes.byref(path)) != 0:
        raise RuntimeError("Não foi possível localizar a Área de Trabalho configurada no Windows.")
    try:
        directory = Path(path.value)
    finally:
        ctypes.windll.ole32.CoTaskMemFree(ctypes.cast(path, ctypes.c_void_p))
    if not directory.is_absolute() or not directory.is_dir():
        raise RuntimeError("A Área de Trabalho não está disponível.")
    return directory


def create_document(title, text, sources=(), directory=None):
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    title = " ".join(str(title).split())[:120]
    if not title or not text.strip() or len(text) > 20000:
        raise ValueError("O documento exige título e até 20.000 caracteres de texto.")
    if any(ord(char) < 32 and char not in "\n\r\t" for char in title + text):
        raise ValueError("O texto contém caracteres de controle não suportados.")
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Cm(2)
    for name in ("Normal", "Title", "Heading 1", "Heading 2"):
        font = document.styles[name].font
        font.name, font.color.rgb, font.underline = "Calibri", RGBColor(0, 0, 0), False
    document.styles["Normal"].font.size = Pt(11)
    document.styles["Normal"].paragraph_format.space_after = Pt(5)
    document.styles["Normal"].paragraph_format.line_spacing = 1.05
    document.styles["Title"].font.size = Pt(16)
    document.add_paragraph(title, "Title")
    for paragraph in text.split("\n\n"):
        if paragraph.strip():
            document.add_paragraph(paragraph.strip())
    if sources:
        document.add_heading("Fontes consultadas", 2)
        for index, source in enumerate(sources, 1):
            url = public_https_url(source["url"])
            paragraph = document.add_paragraph()
            link = OxmlElement("w:hyperlink")
            relation = document.part.relate_to(
                url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
                is_external=True,
            )
            link.set(qn("r:id"), relation)
            run = OxmlElement("w:r")
            properties = OxmlElement("w:rPr")
            color = OxmlElement("w:color")
            color.set(qn("w:val"), "000000")
            properties.append(color)
            run.append(properties)
            label = OxmlElement("w:t")
            label.text = f"[{index}] {str(source.get('title', 'Fonte'))[:90]} — {urlsplit(url).hostname}"
            run.append(label)
            link.append(run)
            paragraph._p.append(link)
    buffer = io.BytesIO()
    document.save(buffer)
    Document(io.BytesIO(buffer.getvalue()))  # Validate package before creating any destination file.
    directory = Path(directory) if directory is not None else desktop_directory()
    stem = re.sub(r"[^\w -]", "", title, flags=re.UNICODE).strip(" .")[:70] or "Relatorio"
    path = directory / f"{stem}-{uuid4().hex[:8]}.docx"
    with path.open("xb") as output:
        output.write(buffer.getvalue())
    return path
