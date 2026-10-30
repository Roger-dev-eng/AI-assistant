import io
import json
import ssl
from dataclasses import replace
from threading import Event
from unittest.mock import Mock, patch

import pytest
from docx import Document

from assistant.ai.bedrock_provider import BedrockProvider
from assistant.cloud.usage_tracker import UsageTracker
from assistant.config import Settings
from assistant.core.intent_router import IntentRouter
from assistant.core.research import ResearchManager, run_research
from assistant.storage.tasks import TaskStore
from assistant.tools.documents import create_document
from assistant.tools.public_pages import PublicConnection, public_addresses, read_article


def test_document_uses_supplied_directory_and_never_overwrites_existing_file(tmp_path):
    source = {"title": "Fonte pública", "url": "https://example.org/article"}
    first = create_document("../Relatório", "Conclusão principal.\n\nExplicação baseada na fonte [1].", [source], tmp_path)
    original = first.read_bytes()
    second = create_document("../Relatório", "Outro relatório.", [source], tmp_path)
    assert first.parent == second.parent == tmp_path
    assert first != second and first.read_bytes() == original
    document = Document(io.BytesIO(original))
    assert document.paragraphs[0].style.name == "Title"
    assert "Conclusão principal" in document.paragraphs[1].text
    assert any(relation.target_ref == source["url"] for relation in document.part.rels.values())


def test_research_has_separate_context_and_only_reports_completion_after_file_exists(tmp_path):
    store = TaskStore(tmp_path / "usage.sqlite3")
    identifier = store.create("Energia solar", "Energia solar", 1)
    browser = Mock()
    browser.read_search_results.return_value = json.dumps({"untrusted_search_results": [
        {"title": "Fonte", "url": "https://example.org/article"},
    ]})
    provider = Mock(client=None, last_stop_reason="end_turn")
    provider.converse.return_value = {"content": [{"text": "Conclusão fundamentada na fonte [1]."}]}
    factory = Mock(return_value=provider)

    def writer(title, text, sources):
        assert store.get(identifier)["status"] == "creating_document"
        return create_document(title, text, sources, tmp_path)

    settings = replace(Settings(), cloud_enabled=True)
    run_research(identifier, store, settings, Event(), browser=browser,
                 reader=lambda url, cancelled: {"url": url, "text": "Conteúdo público da fonte."},
                 writer=writer, provider_factory=factory)
    task = store.get(identifier)
    assert task["status"] == "completed"
    assert Document(task["output_path"]).paragraphs[0].text == "Energia solar"
    assert factory.call_args.kwargs["tool_names"] == ()
    assert factory.call_args.args[0].max_output_tokens > settings.max_output_tokens
    assert settings.max_output_tokens == 512


def test_failed_sources_or_cancellation_do_not_call_model_or_create_document(tmp_path):
    store = TaskStore(tmp_path / "usage.sqlite3")
    identifier = store.create("Assunto", "Relatório", 1)
    browser, provider, writer = Mock(), Mock(), Mock()
    browser.read_search_results.return_value = json.dumps({"untrusted_search_results": [
        {"title": "Fonte", "url": "https://example.org/article"},
    ]})
    run_research(identifier, store, replace(Settings(), cloud_enabled=True), Event(), browser=browser,
                 reader=Mock(side_effect=ValueError("blocked")), writer=writer, provider_factory=provider)
    assert store.get(identifier)["status"] == "failed"
    provider.assert_not_called()
    writer.assert_not_called()
    cancelled = Event()
    cancelled.set()
    run_research(identifier, store, Settings(), cancelled, browser=browser, writer=writer, provider_factory=provider)
    assert store.get(identifier)["status"] == "cancelled"
    provider.assert_not_called()


def test_public_reader_preserves_query_and_blocks_private_redirect_before_connect():
    response = Mock(status=302)
    response.getheader.return_value = "https://127.0.0.1/private"
    connection = Mock()
    connection.getresponse.return_value = response
    factory = Mock(return_value=connection)
    with pytest.raises(ValueError):
        read_article("https://example.org/article?id=1", Event(), factory)
    factory.assert_called_once()
    assert connection.request.call_args.args[1] == "/article?id=1"
    connection.close.assert_called_once()


def test_dns_rebinding_is_prevented_by_connecting_to_verified_ip():
    with patch("assistant.tools.public_pages.socket.getaddrinfo", return_value=[(0, 0, 0, "", ("127.0.0.1", 443))]):
        with pytest.raises(ValueError):
            public_addresses("example.org")
    connection = PublicConnection("example.org", context=ssl.create_default_context())
    connection._context = Mock()
    with (
        patch("assistant.tools.public_pages.public_addresses", return_value=["8.8.8.8"]),
        patch("assistant.tools.public_pages.socket.create_connection") as connect,
    ):
        connection.connect()
    assert connect.call_args.args[0] == ("8.8.8.8", 443)
    assert connection._context.wrap_socket.call_args.kwargs["server_hostname"] == "example.org"


def test_research_model_request_omits_tools_and_uses_shared_budget(tmp_path):
    tracker = UsageTracker(tmp_path / "usage.sqlite3")
    client = Mock()
    client.converse.return_value = {"usage": {"inputTokens": 10, "outputTokens": 10},
                                    "output": {"message": {"content": [{"text": "Report"}]}}, "stopReason": "end_turn"}
    provider = BedrockProvider(replace(Settings(), cloud_enabled=True), tracker, client, tool_names=())
    provider.converse([{"role": "user", "content": [{"text": "Write a short report"}]}])
    assert "toolConfig" not in client.converse.call_args.kwargs
    assert tracker.summary()[0] == 1
    tracker.close()


def test_research_command_is_one_independent_task():
    actions = IntentRouter().route("Pesquise energia solar e gere um relatório de uma página na área de trabalho")
    assert len(actions) == 1
    assert actions[0].name == "start_research_report"
    assert actions[0].arguments["pages"] == 1


def test_resume_research_requires_unfinished_inactive_task_and_preserves_original(tmp_path):
    manager = ResearchManager(tmp_path / "usage.sqlite3")
    identifier = manager.store.create("Energia solar", "Relatório solar", 1)
    manager.store.update(identifier, "failed", message="Interrupção simulada")
    try:
        with patch.object(manager, "submit", return_value="new-task") as submit:
            assert manager.resume(Settings(), identifier) == "new-task"
            assert submit.call_args.args[1:] == ("Energia solar", "Relatório solar", 1)
            assert manager.store.get(identifier)["status"] == "failed"
            with pytest.raises(ValueError, match="nova tentativa"):
                manager.resume(Settings(), identifier)
            manager.restarts.clear()
            manager.jobs[identifier] = (Event(), Mock(done=Mock(return_value=False)))
            with pytest.raises(ValueError, match="execução"):
                manager.resume(Settings(), identifier)
            manager.jobs.clear()
            manager.store.update(identifier, "completed", output_path="existing.docx")
            with pytest.raises(ValueError, match="documento"):
                manager.resume(Settings(), identifier)
            assert submit.call_count == 1
    finally:
        manager.close()
